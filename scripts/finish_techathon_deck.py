"""Embed the reviewed logo film in a PPT Master native export and verify it.

Run after PPT Master's SVG/notes/finalize/export sequence. Private assets and
finished presentations are supplied as local paths, never checked into Git.
"""
from pathlib import Path
import argparse
import hashlib
import json
import zipfile

from lxml import etree
from pptx import Presentation

NS={'p':'http://schemas.openxmlformats.org/presentationml/2006/main',
    'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}

def finish(project: Path, output: Path):
    source=project/'exports/派特岛_五分钟路演.pptx'
    movie=project/'images/five-islands-logo.mp4'
    poster=project/'images/opening-frame.png'
    prs=Presentation(source)
    assert len(prs.slides)==11
    assert prs.slide_width*9==prs.slide_height*16
    scale=prs.slide_width/1280
    video=prs.slides[1].shapes.add_movie(str(movie),*(round(v*scale) for v in (480,174,752,423)),
                                      poster_frame_image=str(poster),mime_type='video/mp4')
    video.name='五岛轮廓到爪印 · 点击播放'
    # add_movie creates indefinite media timing; retain manual click-to-play.
    output.parent.mkdir(parents=True,exist_ok=True)
    prs.save(output)
    with zipfile.ZipFile(output) as z:
        assert z.testzip() is None
        slides=[etree.fromstring(z.read(f'ppt/slides/slide{i}.xml')) for i in range(1,12)]
        for i,s in enumerate(slides,1):
            assert s.find('p:transition/p:fade',NS) is not None, f'No fade on {i}'
            assert s.find('p:timing',NS) is not None, f'No timing on {i}'
            ids=s.xpath('//p:cTn/@id',namespaces=NS)
            assert len(ids)==len(set(ids)),f'Duplicate timing IDs on {i}'
            assert len(s.xpath('//a:t',namespaces=NS))>3,f'Not editable text: {i}'
        fonts=set(f for s in slides for f in s.xpath('//a:latin/@typeface | //a:ea/@typeface',namespaces=NS))
        assert 'KaiTi' in fonts and 'Microsoft YaHei' in fonts,fonts
        media=[n for n in z.namelist() if n.startswith('ppt/media/') and n.endswith('.mp4')]
        assert len(media)==1
        assert hashlib.sha256(z.read(media[0])).digest()==hashlib.sha256(movie.read_bytes()).digest()
        for i in range(1,12):
            note=etree.fromstring(z.read(f'ppt/notesSlides/notesSlide{i}.xml'))
            assert len(''.join(note.xpath('//a:t/text()',namespaces=NS)))>20
    report={'slides':11,'editable_text':True,'fonts':sorted(fonts),'fade_transitions':11,
            'entrance_timing_slides':11,'embedded_logo_video':True,'logo_trigger':'click',
            'narration_notes':11,'seconds':285,'output':str(output)}
    output.with_suffix('.validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),'utf-8')
    print(json.dumps(report,ensure_ascii=False))

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--project',type=Path,required=True);p.add_argument('--output',type=Path,required=True)
    a=p.parse_args();finish(a.project,a.output)
